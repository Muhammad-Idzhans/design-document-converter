"""
design_doc_agents.py — Design Document Generator: AI agents, prompts, and processing pipeline.

Contains everything specific to the Design Document generation workflow:
- PPTX extraction (shapes, images, tables, speaker notes)
- Azure Content Understanding analysis
- GPT-4o Vision image analysis
- Orchestrator prompt (TOC generation)
- Writer prompt (section drafting with Agent + fallback)
- Appendix agent
- Image placeholder post-processing
- Design Document title page and front matter
- Master DOCX builder (convert_md_to_docx)
- Background processing pipelines (upload + generation)
"""

import os
import json
import re
import base64
import time
import shutil
import traceback
import requests
from pathlib import Path
from typing import Optional, Dict, Any, List

from pptx import Presentation

from config import (
    CONTENT_UNDERSTANDING_ENDPOINT, CONTENT_UNDERSTANDING_KEY, API_VERSION,
    AZURE_OPENAI_ENDPOINT, AZURE_OPENAI_KEY, AZURE_OPENAI_DEPLOYMENT_NAME,
    AGENT_OPENAI_ENDPOINT, AGENT_OPENAI_KEY,
    ORCHESTRATOR_DEPLOYMENT, WRITER_DEPLOYMENT, AGENT_ASSISTANT_ID,
    SCRIPT_DIR, UPLOAD_DIR
)
from task_manager import update_task, get_task, upload_task_to_blob
from doc_styles import (
    convert_pptx_to_pdf,
    _pandoc_md_to_raw_docx,
    _apply_header_footer,
    _fix_tables,
    _fix_typography_and_captions,
    _fix_images,
    _strip_heading_numbers,
    _update_toc_with_libreoffice
)


# -----------------------------
# Upload Background Worker
# -----------------------------
def process_initial_upload(task_id: str, file_path: str, logo_path: Optional[str], src_name: str, task_dir: Path):
    try:
        update_task(task_id, {
            "status": "processing_upload",
            "step_name": "Generating Slide Thumbnails...",
            "progress": 20
        })

        # This is the slow part that was causing Vercel to timeout
        generate_slide_thumbnails(file_path, task_dir)

        update_task(task_id, {
            "step_name": "Extracting Slide Data...",
            "progress": 70
        })

        preview_data = extract_preview(file_path, task_id)

        # Mark as completely ready for the Next.js Preview Page
        update_task(task_id, {
            "status": "upload_complete",
            "step_name": "Ready for Preview",
            "progress": 100,
            "filename": src_name,
            "file_path": file_path,
            "logo_path": logo_path,
            "preview_data": preview_data
        })
    except Exception as e:
        update_task(task_id, {
            "status": "failed",
            "step_name": "Upload Failed",
            "error": str(e)
        })


# -----------------------------
# Thumbnails / Preview
# -----------------------------
def generate_slide_thumbnails(pptx_path: str, output_dir: Path) -> None:
    import fitz

    temp_pdf = str(output_dir / "temp_thumbnail_source.pdf")
    convert_pptx_to_pdf(pptx_path, temp_pdf)

    if not os.path.exists(temp_pdf):
        print("Failed to generate PDF for thumbnail extraction.")
        return

    doc = None
    try:
        doc = fitz.open(temp_pdf)
        for i in range(len(doc)):
            page = doc.load_page(i)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            export_path = str(output_dir / f"slide_{i+1}.jpg")
            pix.save(export_path)
    except Exception as e:
        print(f"Failed to generate thumbnails via PyMuPDF: {e}")
    finally:
        try:
            if doc is not None:
                doc.close()
        except Exception as e:
            print(f"[generate_slide_thumbnails] Failed to close PDF document: {e}")
        try:
            if os.path.exists(temp_pdf):
                os.remove(temp_pdf)
        except Exception as e:
            print(f"[generate_slide_thumbnails] Failed to remove temp PDF: {e}")


def extract_preview(pptx_path: str, task_id: str) -> Dict[str, Any]:
    prs = Presentation(pptx_path)
    slides_data = []
    diagram_count = 0

    for idx, slide in enumerate(prs.slides, 1):
        title = ""
        notes = "No presenter note"

        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                notes = n

        for shape in slide.shapes:
            if shape.shape_type == 13:
                diagram_count += 1

        slides_data.append({
            "slide_number": idx,
            "title": title if title else f"Slide {idx}",
            "notes": notes,
            "thumbnail_url": f"/api/thumbnails/{task_id}/slide_{idx}.jpg"
        })

    return {
        "filename": Path(pptx_path).name,
        "total_slides": len(prs.slides),
        "diagrams_detected": diagram_count,
        "slides": slides_data
    }


# -----------------------------
# PPTX extraction
# -----------------------------
def extract_shapes(shapes, slide_info: Dict[str, Any], ctx: Dict[str, Any]) -> None:
    for shape in shapes:
        # GroupShape recursion
        if shape.shape_type == 6:
            extract_shapes(shape.shapes, slide_info, ctx)
            continue

        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text and text != slide_info.get("title"):
                    slide_info["text_content"].append(text)

        if shape.shape_type == 13:
            # Safely try to access shape.image — some PPTX files have linked
            # or broken image references that raise AttributeError.
            try:
                img = shape.image
            except (AttributeError, KeyError, ValueError) as e:
                print(f"[extract_shapes] Skipping broken/linked image on slide {ctx['slide_idx']}: {e}")
                continue
            
            ctx["image_counter"] += 1
            ext = img.content_type.split("/")[-1]


            if ext in ["x-wmf", "wmf", "x-emf", "emf"]:
                ext = "png"

            filename = f"slide_{ctx['slide_idx']:03d}_img_{ctx['image_counter']:03d}.{ext}"
            filepath = ctx["images_dir"] / filename

            if ext in ["png"]:
                try:
                    import io
                    from PIL import Image
                    wmf_img = Image.open(io.BytesIO(img.blob))
                    wmf_img.save(filepath, format="PNG")
                    content_type = "image/png"
                    size_bytes = os.path.getsize(filepath)
                except Exception as e:
                    print(f"[extract_shapes] WMF/EMF to PNG conversion failed, saving raw: {e}")
                    filename = f"slide_{ctx['slide_idx']:03d}_img_{ctx['image_counter']:03d}.x-wmf"
                    filepath = ctx["images_dir"] / filename
                    with open(filepath, "wb") as f:
                        f.write(img.blob)
                    content_type = img.content_type
                    size_bytes = len(img.blob)
            else:
                if ext == "jpeg":
                    ext = "jpg"
                filename = f"slide_{ctx['slide_idx']:03d}_img_{ctx['image_counter']:03d}.{ext}"
                filepath = ctx["images_dir"] / filename
                with open(filepath, "wb") as f:
                    f.write(img.blob)
                content_type = img.content_type
                size_bytes = len(img.blob)

            slide_info["images"].append({
                "image_id": f"img_{ctx['slide_idx']:03d}_{ctx['image_counter']:03d}",
                "filename": filename,
                "file_path": str(filepath),
                "content_type": content_type,
                "size_bytes": size_bytes,
            })

        if shape.has_table:
            table = shape.table
            headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
            rows = []
            for row_idx, row in enumerate(table.rows):
                if row_idx == 0:
                    continue
                row_data = [cell.text.strip() for cell in row.cells]
                rows.append(row_data)

            slide_info["tables"].append({
                "headers": headers,
                "rows": rows,
                "total_rows": len(table.rows),
                "total_cols": len(table.columns),
            })


def extract_with_pptx(pptx_path: str, output_dir: Path) -> Dict[str, Any]:
    prs = Presentation(pptx_path)
    images_dir = output_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    slides_data = []
    ctx = {"image_counter": 0, "images_dir": images_dir, "slide_idx": 0}

    for slide_idx, slide in enumerate(prs.slides, 1):
        ctx["slide_idx"] = slide_idx
        slide_info = {
            "slide_number": slide_idx,
            "title": None,
            "text_content": [],
            "speaker_notes": None,
            "images": [],
            "tables": []
        }

        if slide.shapes.title:
            slide_info["title"] = slide.shapes.title.text.strip()

        extract_shapes(slide.shapes, slide_info, ctx)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_info["speaker_notes"] = notes

        slides_data.append(slide_info)

    return {
        "metadata": {
            "source_file": Path(pptx_path).name,
            "total_slides": len(prs.slides),
            "total_images_extracted": ctx["image_counter"],
            "slides_with_notes": sum(1 for s in slides_data if s["speaker_notes"]),
        },
        "slides": slides_data,
    }


# -----------------------------
# Content Understanding
# -----------------------------
def analyze_with_content_understanding(pdf_path: str) -> Optional[Dict[str, Any]]:
    if not CONTENT_UNDERSTANDING_ENDPOINT or not CONTENT_UNDERSTANDING_KEY:
        return None

    analyze_binary_url = (
        f"{CONTENT_UNDERSTANDING_ENDPOINT}/contentunderstanding/analyzers/"
        f"design_document_extractor:analyzeBinary?api-version={API_VERSION}"
    )

    with open(pdf_path, "rb") as f:
        file_bytes = f.read()

    headers = {
        "Ocp-Apim-Subscription-Key": CONTENT_UNDERSTANDING_KEY,
        "Content-Type": "application/pdf",
    }

    response = requests.post(analyze_binary_url, headers=headers, data=file_bytes)
    if response.status_code not in (200, 202):
        return _analyze_with_base64_fallback(pdf_path)

    operation_url = response.headers.get("Operation-Location")
    if not operation_url:
        if response.status_code == 200:
            return response.json()
        return None

    return _poll_for_result(operation_url)


def _analyze_with_base64_fallback(pdf_path: str) -> Optional[Dict[str, Any]]:
    analyze_url = (
        f"{CONTENT_UNDERSTANDING_ENDPOINT}/contentunderstanding/analyzers/"
        f"design_document_extractor:analyze?api-version={API_VERSION}"
    )

    with open(pdf_path, "rb") as f:
        file_bytes = f.read()

    b64_content = base64.b64encode(file_bytes).decode("utf-8")
    data_uri = f"data:application/pdf;base64,{b64_content}"

    headers = {
        "Ocp-Apim-Subscription-Key": CONTENT_UNDERSTANDING_KEY,
        "Content-Type": "application/json",
    }

    payload = {"inputs": [{"data": data_uri}]}
    response = requests.post(analyze_url, headers=headers, json=payload)

    if response.status_code not in (200, 202):
        return None

    operation_url = response.headers.get("Operation-Location")
    if operation_url:
        return _poll_for_result(operation_url)

    return response.json() if response.status_code == 200 else None


def _poll_for_result(operation_url: str, max_retries: int = 60, interval: int = 3) -> Optional[Dict[str, Any]]:
    headers = {"Ocp-Apim-Subscription-Key": CONTENT_UNDERSTANDING_KEY}
    for _ in range(max_retries):
        time.sleep(interval)
        response = requests.get(operation_url, headers=headers)
        if response.status_code != 200:
            continue
        result = response.json()
        status = result.get("status", "")
        if status == "Succeeded":
            return result
        if status in ("Failed", "Cancelled"):
            return None
    return None


# -----------------------------
# Vision Image Analysis
# -----------------------------
def analyze_images_with_vision(task_id: str, slides_data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not AZURE_OPENAI_ENDPOINT or not AZURE_OPENAI_KEY:
        return slides_data

    base_endpoint = AZURE_OPENAI_ENDPOINT.split("/openai")[0].rstrip("/")
    url = f"{base_endpoint}/openai/v1/chat/completions"
    headers = {
        "api-key": AZURE_OPENAI_KEY,
        "Authorization": f"Bearer {AZURE_OPENAI_KEY}",
        "Content-Type": "application/json"
    }

    vision_prompt = (
        "You are an Expert Azure Solutions Architect. Analyze this image extracted from a pre-sales presentation.\n"
        "1. If this image is a company logo, a decorative background, a transition slide, or a generic stock photo, reply ONLY with the word: DECORATIVE.\n"
        "2. If this image is a technical architecture diagram, a data workflow, or a system component map, provide a highly detailed, professional text description of the data flow, the components involved, and how they connect.\n"
        "Do not include conversational filler."
    )

    for slide in slides_data:
        for img in slide.get("images", []):
            filepath = img.get("file_path")
            if not filepath or not os.path.exists(filepath):
                continue

            with open(filepath, "rb") as f:
                base64_image = base64.b64encode(f.read()).decode("utf-8")

            mime_type = img.get("content_type", "image/png")

            payload = {
                "model": AZURE_OPENAI_DEPLOYMENT_NAME,
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": vision_prompt},
                            {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{base64_image}"}}
                        ]
                    }
                ],
                "max_tokens": 800,
                "temperature": 0.0
            }

            try:
                response = requests.post(url, headers=headers, json=payload, timeout=120)
                if response.status_code == 200:
                    data = response.json()
                    description = data["choices"][0]["message"]["content"].strip()
                    img["ai_description"] = description
                    
                    # Accumulate token usage
                    usage = data.get("usage", {})
                    task = get_task(task_id)
                    if task and "cost_metrics" in task:
                        task["cost_metrics"]["vision_tokens_prompt"] += usage.get("prompt_tokens", 0)
                        task["cost_metrics"]["vision_tokens_completion"] += usage.get("completion_tokens", 0)
                        update_task(task_id, {"cost_metrics": task["cost_metrics"]})
                        
            except Exception as e:
                print(f"[analyze_images_with_vision] Vision API call failed for {img.get('filename', 'unknown')}: {e}")

    return slides_data


def merge_extraction_results(pptx_data: Dict[str, Any], cu_data: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    result = {
        "metadata": pptx_data["metadata"],
        "extraction_sources": ["python-pptx", "azure-openai-vision"],
        "slides": pptx_data["slides"],
        "content_understanding": None,
    }

    if cu_data:
        result["extraction_sources"].append("content-understanding")
        cu_result = cu_data.get("result", cu_data)
        contents = cu_result.get("contents", [])
        if contents:
            markdown_content = contents[0].get("markdown", "")
            result["content_understanding"] = {
                "markdown": markdown_content,
                "markdown_length": len(markdown_content),
            }

    return result


# -----------------------------
# Orchestrator & Writer
# -----------------------------
def generate_table_of_contents(task_id: str, extraction_payload: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if not AGENT_OPENAI_ENDPOINT or not AGENT_OPENAI_KEY:
        return None

    base_endpoint = AGENT_OPENAI_ENDPOINT.split("/openai")[0].rstrip("/")
    url = f"{base_endpoint}/openai/v1/chat/completions"
    headers = {
        "api-key": AGENT_OPENAI_KEY,
        "Authorization": f"Bearer {AGENT_OPENAI_KEY}",
        "Content-Type": "application/json"
    }

    orchestrator_prompt = (
        "You are an Expert Enterprise Solutions Architect from Enfrasys Solutions. Your task is to analyze an extracted solution design presentation "
        "(provided as a combined JSON payload) and produce a detailed, bespoke Table of Contents (Outline) for a formal Microsoft Word Solution Design Document.\n\n"
        "INSTRUCTIONS:\n"
        "1. Analyze the payload deeply: read slide text, speaker notes, and diagram descriptions to fully understand the technical context, the scope of the project, and the design decisions made.\n"
        "2. DYNAMIC MAPPING (CRITICAL): Do NOT blindly copy a static template. You must build the Table of Contents based ONLY on the actual technologies and concepts present in the payload. If the slides do not mention Virtual Machines, do not create a VM section. If they do not mention NSGs, do not create an NSG section.\n"
        "3. REQUIRED SECTION FLOW — Follow this logical progression, adapting the specific section titles to the actual content:\n"
        "   - 2.0 Executive Summary (CRITICAL: This section MUST contain ONLY exactly three sub-sections: Project Overview, Document Purpose, and Document Audience. Do NOT generate any other sub-sections or stakeholder tables here).\n"
        "   - 3.0 Discovery & Risk Assessment (CONDITIONAL — Read these rules carefully):\n"
        "       DETECTION: Search the slide data for slides titled 'Assessment Outcome', 'Assessment Outcomes', 'Current State', 'Risk and Mitigation', 'Design Constraints', or slides with titles starting with 'Risk #', OR slides containing a multi-row risk/constraint table.\n"
        "       SCENARIO A (Both assessment AND risk slides found): Create section '3.0 Discovery & Risk Assessment' with two sub-sections:\n"
        "           - 3.1 Assessment Outcome: Instruct the Writer to summarize the current environment assessment findings, including any inventory/discovery tables from the slides.\n"
        "           - 3.2 Risk and Mitigation: BEFORE writing this sub-section, detect which RISK FORMAT the slides use:\n"
        "\n"
        "               FORMAT A — Per-Risk Slides: If the slides contain MULTIPLE separate slides each titled 'Risk #1', 'Risk #2', etc., with their own 2-column Problem/Impact/Action/By table → instruct the Writer to document each risk as its own H3 sub-sub-section (### Risk #1 Title, ### Risk #2 Title). For EACH risk, generate a horizontal Markdown table with this EXACT format:\n"
        "                   | Field | Detail |\n"
        "                   |---|---|\n"
        "                   | Problem Description | [actual detail from slides] |\n"
        "                   | Impact if Not Remediated | [actual detail from slides] |\n"
        "                   | Required Action | [actual detail from slides] |\n"
        "                   | Action By | [actual entity from slides] |\n"
        "\n"
        "               FORMAT B — Consolidated Risk Table: If the slides contain ONE SINGLE slide with a multi-row risk/constraint table (e.g., columns like 'No', 'Constraint/Risk', 'Risk Level', 'Approach/Mitigation') → instruct the Writer to produce ONE top-header Markdown table that preserves the ORIGINAL PPTX column headers EXACTLY. Do NOT force the Problem/Impact/Action/By format. Do NOT split into per-risk sub-sections. Example:\n"
        "                   | No | Constraint / Risk Identified | Risk Level | Approach / Mitigation |\n"
        "                   |---|---|---|---|\n"
        "                   | 1 | [from slides] | [from slides] | [from slides] |\n"
        "                   | 2 | [from slides] | [from slides] | [from slides] |\n"
        "\n"
        "               FORMAT C — Prose/Bullets Only: If risks are described only in bullet points or prose (no table) → instruct the Writer to output a numbered list with 1-2 sentences of consulting expansion for each risk.\n"
        "\n"
        "           CRITICAL TABLE SYNTAX RULES (apply to whichever format is chosen):\n"
        "               - The header row AND the '|---|---|' separator row are MANDATORY. Without them, the table will render as plain text.\n"
        "               - Every row MUST have leading and trailing pipes (`|`).\n"
        "               - Do NOT mix formats. Pick ONE format based on the slide data and stick with it.\n"
        "       SCENARIO B (ONLY assessment slides found, NO risk slides): Create section '3.0 Assessment Outcome' with NO sub-sections. Instruct the Writer to write all assessment content directly under this heading.\n"
        "       SCENARIO C (ONLY risk slides found, NO assessment slides): Create section '3.0 Risk and Mitigation' with NO sub-sections. Instruct the Writer to write all risk content directly under this heading.\n"
        "       SCENARIO D (NEITHER assessment NOR risk slides found): SKIP this section entirely. Start the next section at 3.0 instead.\n"
        "       All subsequent section numbers MUST shift accordingly based on whether this section exists or not.\n"
        "   - Next: Technology Overview (Instruct the Writer to define the core technologies mentioned in the slides, e.g., Fabric, M365, Entra ID). Number this section sequentially after the previous section.\n"
        "   - Next: Network & Security Design (Include Sub-sections for Connectivity, Gateways, or NSGs ONLY if present in the data)\n"
        "   - Next: Roles and Access (Include Sub-sections for Admin Roles, Conditional Access, or Workspace Mapping ONLY if present)\n"
        "   - Next: Platform Design & Decisions (Include Sub-sections for Workflows, Resource Organization, or Data Governance ONLY if present)\n"
        "       SPECIAL HANDLING FOR MULTI-PART SETTINGS TABLES: If the slides contain multiple parts of the same settings table (e.g., 'Settings Part 1/4', 'Settings Part 2/4', 'Settings Part 3/4', 'Settings Part 4/4', or 'Meetings Settings 1/7' through 'Meetings Settings 7/7'), you MUST combine all parts into a SINGLE sub-section. Do NOT create separate sub-sections for each part. For example, if there are 7 parts of Meetings Settings, create only ONE sub-section titled '6.X Meetings Settings' (do NOT include 'Part X of Y' in the title) and instruct the Writer to combine the rows from all parts into one consolidated markdown table. In the generation_instructions for this sub-section, explicitly state: 'Combine the policy/settings tables from these slides into one single table. Do not paraphrase or summarize. Apply the FAITHFUL TRANSCRIPTION RULE.'\n"
        "   - Next: Deployment & Migration Approach (Include ONLY if the slides contain migration timelines, phases, or strategies)\n"
        "   - DO NOT INCLUDE AN APPENDIX SECTION. The Appendix will be generated separately by a dedicated agent after the main content is complete. Stop the TOC at the last main content section (Deployment & Migration Approach or whichever is last).\n"
        "4. Output the outline strictly as a JSON object matching this schema:\n"
        "{\n"
        ' "client_name": "string (Exact client name from title slide)",\n'
        ' "client_name_full": "string | null",\n'
        ' "project_title": "string",\n'
        ' "sections": [\n'
        "  {\n"
        '   "section_number": "string (e.g., 2.0, 4.1)",\n'
        '   "section_title": "string",\n'
        '   "mapped_slides": [int],\n'
        '   "generation_instructions": "string (List the exact required tables WITH column names if applicable, explicit instructions to add boilerplate definitions, and instructions to assign Enfrasys vs. Client responsibilities.)"\n'
        "  }\n"
        " ]\n"
        "}\n"
    )
    
    payload = {
        "model": ORCHESTRATOR_DEPLOYMENT,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": orchestrator_prompt},
            {"role": "user", "content": json.dumps(extraction_payload, default=str)}
        ],
        "temperature": 0.2,
        "max_tokens": 16384
    }

    try:
        response = requests.post(url, headers=headers, json=payload, timeout=180)
        if response.status_code == 200:
            data = response.json()
            content = data["choices"][0]["message"]["content"]
            
            # Accumulate token usage
            usage = data.get("usage", {})
            task = get_task(task_id)
            if task and "cost_metrics" in task:
                task["cost_metrics"]["llm_tokens_prompt"] += usage.get("prompt_tokens", 0)
                task["cost_metrics"]["llm_tokens_completion"] += usage.get("completion_tokens", 0)
                update_task(task_id, {"cost_metrics": task["cost_metrics"]})
                
            return json.loads(content)
        return None
    except Exception as e:
        print(f"[generate_table_of_contents] Failed to parse TOC response: {e}")
        return None


def write_document_sections(task_id: str, toc: Dict[str, Any], extraction_payload: Dict[str, Any]) -> str:
    base_endpoint = AGENT_OPENAI_ENDPOINT.split("/openai")[0].rstrip("/")
    url = f"{base_endpoint}/openai/v1/chat/completions"
    headers = {
        "api-key": AGENT_OPENAI_KEY,
        "Authorization": f"Bearer {AGENT_OPENAI_KEY}",
        "Content-Type": "application/json"
    }

    final_document_markdown = f"# {toc.get('client_name', '')} {toc.get('project_title', 'Solution Design Document')}".strip() + "\n\n"

    writer_system_prompt = """
        You are a Lead Enterprise Architect from Enfrasys Solutions writing a formal Solution Design Document. Your output will be converted directly to a Microsoft Word document for a client.
        Write in authoritative, professional language ("Enfrasys recommends", "[CLIENT_NAME] shall implement"). Be specific — reference the actual client name, technology components, and design decisions from the provided slide data.
        CRITICAL: Do NOT wrap your entire response in a ```markdown code block. Just return the raw text.

        -- CONTENT DEPTH AND EXPANSION RULES (CRITICAL FOR LENGTH & QUALITY) --
        1. PROSE BEFORE TABLES: Human consultants do not just output naked tables. For every H2 and H3 sub-section, you MUST write 2-3 comprehensive paragraphs explaining the architecture, the "why" behind the design decisions, and how it specifically benefits the client BEFORE presenting any Markdown table.
        2. EXPAND ON SLIDE DATA: The input provided to you is extracted from PowerPoint slides. Treat these slide bullets as a brief outline. You MUST expand on these points using your expert Enterprise Architect knowledge. Detail the workflows, the security implications, and the operational benefits to make the document feel authoritative and complete.
        3. VISUAL CONTEXT: Whenever you insert an image/diagram, you MUST write a detailed paragraph immediately preceding the image explaining what the diagram represents and how the data flows through it.

        -- PLACEHOLDER PREVENTION RULES (CRITICAL) --
        1. ZERO PLACEHOLDERS: You are STRICTLY FORBIDDEN from outputting any placeholders or template markers (e.g., [CLIENT_NAME], [specific purpose], [Specific client tasks]). If the specific information is not available in the slide data, you MUST rewrite the sentence entirely to be professional, generic, and grammatically complete without a marker (e.g., use "the client" or "the organization").
        2. RECHECK STEP: Before outputting any text, search your draft internally for bracketed placeholders like "[". If you find any (that are not explicitly part of a Markdown URL syntax), you MUST rewrite that content immediately.

        -- CONSULTING EXPANSION RULES --
        1. STATIC BOILERPLATE: Whenever you introduce a major Microsoft technology (e.g., Microsoft Fabric, Exchange Online, Entra ID, Power BI), you MUST provide a standard, formal definition of that technology before discussing the client's specific design. Do not assume the reader knows what the technology does.
        2. RESPONSIBILITY ASSIGNMENT: In any section discussing tasks, migrations, or rollouts, you MUST explicitly assign ownership. Use "Enfrasys Solutions" for vendor tasks and the specific client name (or "the client" if unavailable) for client tasks.
        3. TIMELINES (If Applicable): If the slides mention migration waves or rollouts, automatically structure them into formal consulting timelines (e.g., Pre-Migration (D-7), Cutover (D0), Post-Migration Support (D+1)).

        -- MARKDOWN HEADING RULES --
        - Main Sections MUST use `#` (e.g., `# 5.0 Fabric Design and Decision`).
        - Sub-sections MUST use `##` (e.g., `## 5.1 Data Workflow Design Considerations`).
        - Sub-sub-sections MUST use `###` (e.g., `### 5.7.1 Azure Management Group`).
        - The H1 main section opening paragraph must be 1-2 sentences only. Detail goes in H2/H3 sub-sections.

        -- MANDATORY EXECUTION WORKFLOW --
        You possess the Web Search tool. You MUST execute your task in the following strict sequence:
        1. SEARCH: You must use your Web Search tool to query MS Learn and official Microsoft forums for the latest limits, capabilities, and naming conventions for the technology mentioned.
        IMPORTANT: Do NOT paste URLs or references in the main body. Only store the results for the Appendix.
        2. DRAFT: Write the document section based on both the slide data and the live search results, ensuring heavy prose expansion and removing all placeholders.
        3. RECHECK: Are the technical specifications up to date based on your search? Are there any bracketed placeholders left?
        4. OUTPUT: Proceed to output the final, corrected text.

        -- TABLE SPAM PREVENTION & SCHEMA (CRITICAL) --
        1. DEFAULT TO BULLET POINTS: You MUST default to using standard prose and bullet points for lists of features, considerations, or technology descriptions.
        2. RESTRICTED TABLE USAGE: You are STRICTLY FORBIDDEN from generating a Markdown table unless the slide data is explicitly numeric, a strict matrix, or a migration schedule.
        3. MANDATORY EXCEPTION — POLICY & SETTINGS TABLES: When the source slide data contains a structured configuration table with columns matching ANY of these patterns:
            - "Policy Setting | Recommendation | Description"
            - "Setting | Parameter | Description"
            - "Policy Setting | Parameter | Description / Reason"
            - "Setting | Parameter / Recommendation | Description / Remarks"
            - Any combination of these three columns (Setting/Policy + Value/Recommendation + Description/Remarks)
            You MUST reproduce the ENTIRE table with ALL rows preserved. This is a MANDATORY EXCEPTION to the bullet-point default rule. These are technical configuration deliverables and must NOT be summarized into prose or bullet points.
        4. PANDOC FORMATTING (CRITICAL): You MUST leave a completely blank empty line immediately BEFORE and AFTER every single Markdown table so the document parser renders it correctly as a grid.
        5. TABLE SYNTAX (CRITICAL): You MUST include leading and trailing pipes (`|`) on EVERY row of the table. Example: `| Col A | Col B |`. DO NOT omit the outer edge pipes, or the parser will fail to render the grid!
        6. NO LINE BREAKS IN CELLS (CRITICAL): You are STRICTLY FORBIDDEN from using line breaks (\\n) or paragraph returns inside a table cell. All text within a single cell MUST be written as one continuous line of text, no matter how long it is.
        7. If you do encounter strict data that requires a table, you must use these columns:
            - Design Decisions: | No. | Design Decision | Decision |
            - Task/Migration Rollout: | Item | Activities | Action By | Status |
            - Workspace/Admin capabilities: | Permission | Admin | Contributor | Member | Viewer |
            - Access decisions: | No. | Policy | Decision |
            - VM specifications: | Component | Specification |
            - NSG rules: | Name | Priority | Source | Source Ports | Destination | Destination Ports | Protocol | Access |


        -- FAITHFUL TRANSCRIPTION RULE FOR POLICY/SETTINGS TABLES (CRITICAL) --
        1. COPY-PASTE FIDELITY: For policy/settings configuration tables (as defined in the Mandatory Exception above), you MUST reproduce the table contents EXACTLY as written in the source slides. Do NOT paraphrase, summarize, rewrite, or "improve" the wording.
        2. PRESERVE EVERY ROW: Every row from the source table MUST appear in the output. Do NOT collapse multiple rows into a summary line. If the source has 30 rows, your output MUST have 30 rows.
        3. PRESERVE ALL VALUES VERBATIM:
        - Setting names → copy exactly (e.g., "Bulk email spam action", not "Bulk Email Action")
        - Recommendation values → copy exactly (e.g., "On", "Off", "7 Days", "1 hour expired, notify user 10 mins before signed them out")
        - Description text → copy word-for-word from the source slide
        4. PRESERVE STRUCTURE: If the source slide uses category headers within the table (e.g., "External sharing", "File and folder links", "Unmanaged Device"), preserve these as sub-headings or first-column markers in the output.
        5. NO PARAPHRASING ALLOWED: You are FORBIDDEN from rewording descriptions to "make them sound more professional." The source wording IS the standard.
        6. NO CREATIVE ADDITIONS: Do NOT add your own settings, parameters, or descriptions that are not in the source slides.
        7. MULTI-PART TABLES: If the source contains multiple parts of the same table (e.g., "Part 1/4", "Part 2/4"), you MUST combine all rows from all parts into a SINGLE continuous table under ONE sub-heading. Remove any "Part X of Y" from the title. The final table should contain all rows from Part 1, followed by Part 2, etc.

        NOTE: This rule applies ONLY to policy/settings configuration tables. For all other content (Risk descriptions, Assessment narratives, Technology Overview, Design Decisions, Migration timelines), continue using your default consulting-writing style.


        -- STRICT IMAGE RULES (CRITICAL — NEW SYNTAX) --
        You will be given a list of AVAILABLE IMAGES for this section in the user prompt.
        The system has ALREADY filtered these images to prevent duplication — trust the list completely.

        CASE A — AVAILABLE IMAGES LIST HAS ENTRIES:
        - You MAY insert relevant diagrams using this EXACT placeholder syntax:
            <<IMAGE:filename>>
          Example: <<IMAGE:slide_020_img_003.png>>
        - DO NOT use Markdown image syntax like `path`. The code will handle the actual image embedding.
        - DO NOT invent filenames. Only use filenames EXACTLY from the AVAILABLE IMAGES list.
        - Select only diagrams that add value: architecture diagrams, data flow diagrams, 
          network topologies, or access/security diagrams. Skip decorative or redundant images.
        - CAPTION RULE: After the placeholder, leave a blank empty line, then write a brief 
          caption sentence starting with "Figure:".
          Example structure:
            The diagram below illustrates the data flow from source to destination.
            
            <<IMAGE:slide_020_img_003.png>>
            
            Figure: High-level architecture of the data pipeline.

        CASE B — AVAILABLE IMAGES LIST IS "None" OR EMPTY (CRITICAL):
        You are STRICTLY FORBIDDEN from doing ANY of the following:
        - Writing "The diagram below...", "The following diagram...", "As shown in the diagram..."
        - Writing "The architecture diagram illustrates...", "The figure demonstrates..."
        - Writing any sentence that references a visual, diagram, figure, illustration, or chart
        - Writing any "Figure:" caption
        - Inserting any <<IMAGE:>> placeholder
        - Referencing visual elements from the slide data (even if the slides mention them)
        
        Instead, explain the concept in pure prose. If the slide data mentions "a diagram shows X",
        you MUST rewrite it as "X works as follows..." or "The architecture consists of..." 
        WITHOUT any visual reference.

        REMEMBER: The AVAILABLE IMAGES list is authoritative. If it says None, there is NO image
        for this section — do not fabricate visual references.


        -- NO-REPETITION RULES --
        1. Each table must appear EXACTLY ONCE across the full document. If a table was in a main section, the Appendix must NOT repeat it.
        2. NSG rules belong ONLY in the Security Design section. VM specifications belong ONLY in the Platform Design section.
        3. NO INTERNAL APPENDIX POINTERS: Do not write any sentences that direct the reader to the appendix, such as "Note: Please refer to the appendix..." 
    """
    
    use_agent = False
    openai_client = None
    agent_name = ""
    agent_version = "1"

    if AGENT_ASSISTANT_ID:
        try:
            from azure.identity import DefaultAzureCredential
            from azure.ai.projects import AIProjectClient

            project_client = AIProjectClient(
                endpoint=AGENT_OPENAI_ENDPOINT,
                credential=DefaultAzureCredential()
            )
            openai_client = project_client.get_openai_client()
            agent_parts = AGENT_ASSISTANT_ID.split(":")
            agent_name = agent_parts[0]
            agent_version = agent_parts[1] if len(agent_parts) > 1 else "1"
            use_agent = True
        except Exception as e:
            print(f"Failed to intialize Azure AI Projects Client: {e}")
            use_agent = False

    # Track which images have already been offered to earlier sections
    # This prevents the same image from being offered to multiple sections
    globally_offered_images = set()

    for section in toc.get("sections", []):
        sec_num = section.get("section_number")
        sec_title = section.get("section_title")
        instructions = section.get("generation_instructions")
        mapped_slides = section.get("mapped_slides", [])

        relevant_slide_data = [
            s for s in extraction_payload.get("slides", [])
            if s.get("slide_number") in mapped_slides
        ]

        # Build list of AVAILABLE IMAGES for this section (only non-decorative)
        # AND only images that haven't already been offered to earlier sections
        available_images = []
        for slide in relevant_slide_data:
            for img in slide.get("images", []):
                filename = img.get("filename")
                desc = img.get("ai_description", "")
                
                # Skip if decorative
                if not desc or desc.strip().upper() == "DECORATIVE":
                    continue
                
                # Skip if this image was already offered to an earlier section
                if filename in globally_offered_images:
                    continue
                
                # Offer this image to the current section and mark as used
                available_images.append({
                    "filename": filename,
                    "description": desc[:200]  # Truncate long descriptions
                })
                globally_offered_images.add(filename)

        # Format available images as a readable list for the AI
        if available_images:
            images_list_str = "\n".join([
                f"- {img['filename']} (Description: {img['description']})"
                for img in available_images
            ])
            images_section = f"\nAVAILABLE IMAGES FOR THIS SECTION:\n{images_list_str}\n\nIf relevant, use the placeholder syntax <<IMAGE:filename>> to insert these images.\n"
        else:
            images_section = "\nAVAILABLE IMAGES FOR THIS SECTION: None\n"

        client_name = toc.get("client_name", "")
        client_name_full = toc.get("client_name_full", "")

        context_str = f"Client Acronym: {client_name}\n"
        if client_name_full:
            context_str += f"Client Full Name: {client_name_full} (Use this full name in the content when appropriate)\n"

        user_prompt = (
            f"Write section '{sec_num} {sec_title}' for the design document.\n"
            f"{context_str}\n"
            f"Instructions from Orchestrator: {instructions}\n"
            f"{images_section}\n"
            f"Here is the raw slide data (text, speaker notes, and diagram descriptions):\n"
            f"{json.dumps(relevant_slide_data, default=str)}"
        )

        drafted_text = None

        if use_agent and openai_client:
            try:
                response = openai_client.responses.create(
                    input=[
                        {"type": "message", "role": "system", "content": writer_system_prompt},
                        {"type": "message", "role": "user", "content": user_prompt}
                    ],
                    extra_body={
                        "agent_reference": {
                            "name": agent_name,
                            "version": agent_version,
                            "type": "agent_reference"
                        }
                    }
                )
                drafted_text = response.output_text
                
                # Capture Foundry agent token usage
                try:
                    usage = response.usage
                    task = get_task(task_id)
                    if task and "cost_metrics" in task:
                        # Foundry responses API uses input_tokens / output_tokens
                        prompt_tokens = getattr(usage, 'input_tokens', 0) or getattr(usage, 'prompt_tokens', 0)
                        completion_tokens = getattr(usage, 'output_tokens', 0) or getattr(usage, 'completion_tokens', 0)
                        task["cost_metrics"]["llm_tokens_prompt"] += prompt_tokens
                        task["cost_metrics"]["llm_tokens_completion"] += completion_tokens
                        update_task(task_id, {"cost_metrics": task["cost_metrics"]})
                except Exception as e:
                    print(f"[write_document_sections] Agent usage token extraction failed for '{sec_title}': {e}")
            except Exception as e:
                print(f"[write_document_sections] Agent section drafting failed for '{sec_title}': {e}")

        if not use_agent or not drafted_text:
            payload = {
                "model": WRITER_DEPLOYMENT,
                "messages": [
                    {"role": "system", "content": writer_system_prompt},
                    {"role": "user", "content": user_prompt}
                ],
                "temperature": 0.4,
                "max_tokens": 16384
            }
            try:
                response = requests.post(url, headers=headers, json=payload, timeout=240)
                if response.status_code == 200:
                    data = response.json()
                    drafted_text = data["choices"][0]["message"]["content"].strip()
                    
                    # Accumulate token usage (fallback API)
                    usage = data.get("usage", {})
                    task = get_task(task_id)
                    if task and "cost_metrics" in task:
                        task["cost_metrics"]["llm_tokens_prompt"] += usage.get("prompt_tokens", 0)
                        task["cost_metrics"]["llm_tokens_completion"] += usage.get("completion_tokens", 0)
                        update_task(task_id, {"cost_metrics": task["cost_metrics"]})
            except Exception as e:
                print(f"[write_document_sections] Fallback LLM call failed for '{sec_title}': {e}")

        if drafted_text:
            if drafted_text.startswith("```"):
                drafted_text = drafted_text.split("\n", 1)[-1]
            if drafted_text.endswith("```"):
                drafted_text = drafted_text.rsplit("\n", 1)[0]
            drafted_text = drafted_text.strip()
            final_document_markdown += f"{drafted_text}\n\n"
        else:
            final_document_markdown += f"## {sec_num} {sec_title}\n\n*(Error generating section)*\n\n"

    return final_document_markdown

# -----------------------------
# Appendix Agent (Dedicated)
# -----------------------------
def generate_appendix_from_markdown(task_id: str, document_markdown: str, 
                                     last_section_number: int) -> str:
    """Generate the Appendix by analyzing the completed document markdown.
    
    This runs as a separate agent AFTER the main content is written.
    It reads the full document body, identifies the topics discussed,
    and generates a relevant Appendix with Microsoft Learn references.
    
    Args:
        task_id: Task ID for cost tracking
        document_markdown: The full document markdown (Sections 2.0 - X.0)
        last_section_number: The last section number from the TOC (e.g., 9)
    
    Returns:
        Appendix markdown (just the Appendix part, ready to append)
        Returns empty string if generation fails (safe fallback).
    """
    if not AGENT_OPENAI_ENDPOINT or not AGENT_OPENAI_KEY:
        return ""

    base_endpoint = AGENT_OPENAI_ENDPOINT.split("/openai")[0].rstrip("/")
    url = f"{base_endpoint}/openai/v1/chat/completions"
    headers = {
        "api-key": AGENT_OPENAI_KEY,
        "Authorization": f"Bearer {AGENT_OPENAI_KEY}",
        "Content-Type": "application/json"
    }

    appendix_number = last_section_number + 1
    
    appendix_system_prompt = f"""You are a Microsoft documentation specialist. Your only job is to write the Appendix section for a Solution Design Document.

        You will receive the FULL TEXT of the document body. Read it carefully to identify the specific Microsoft technologies, settings, and topics discussed.

        Then produce ONLY the Appendix section in Markdown format using this EXACT structure:

        # {appendix_number}.0 Appendix

        ## [Topic Area Name]

        Title of Reference: [https://learn.microsoft.com/specific-page](https://learn.microsoft.com/specific-page)

        Title of Reference: [https://learn.microsoft.com/specific-page](https://learn.microsoft.com/specific-page)

        ## [Another Topic Area Name]

        Title of Reference: [https://learn.microsoft.com/specific-page](https://learn.microsoft.com/specific-page)

        RULES:
        1. Generate 2 to 5 H2 sub-sections (use ## heading level). Choose names that describe the actual content areas discussed (examples: "SharePoint Online & OneDrive Configuration", "Microsoft Teams Policies", "Exchange Online Hybrid Mail Flow", "Microsoft Fabric Workspace Management").

        2. Do NOT use generic names like "Introduction/Prerequisites", "Limits and Boundaries", or "Others".

        3. All URLs must be from learn.microsoft.com — no other sources, no third-party blogs, no GitHub.

        4. URLs should point to specific Microsoft Learn pages, not homepages. For example, prefer "learn.microsoft.com/en-us/sharepoint/external-sharing-overview" over "learn.microsoft.com/en-us/sharepoint/".

        5. Each reference should relate to a specific topic discussed in the document.

        6. Leave a blank line between each reference link.

        7. No bullet points, no introductions, no conclusions. Just the headers and references.

        8. Output ONLY the Appendix markdown. Do not wrap in code blocks. Do not add any other text before or after.

        9. CRITICAL LINK FORMAT: Each URL MUST be wrapped in Markdown link syntax `[url](url)` so it renders as a proper clickable hyperlink in Word (blue color, underlined). Do NOT output raw URLs as plain text.
            - WRONG: `Title: https://learn.microsoft.com/page`
            - RIGHT: `Title: [https://learn.microsoft.com/page](https://learn.microsoft.com/page)`
    """

    user_prompt = f"""Here is the document body. Read it and generate the Appendix:

        ---DOCUMENT BODY START---
        {document_markdown}
        ---DOCUMENT BODY END---

        Now generate the Appendix section following the rules in the system prompt.
    """

    payload = {
        "model": WRITER_DEPLOYMENT,
        "messages": [
            {"role": "system", "content": appendix_system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        "temperature": 0.3,
        "max_tokens": 4096
    }

    try:
        response = requests.post(url, headers=headers, json=payload, timeout=180)
        if response.status_code == 200:
            data = response.json()
            appendix_text = data["choices"][0]["message"]["content"].strip()
            
            # Strip code fences if present
            if appendix_text.startswith("```"):
                appendix_text = appendix_text.split("\n", 1)[-1]
            if appendix_text.endswith("```"):
                appendix_text = appendix_text.rsplit("\n", 1)[0]
            
            # Track tokens for cost
            usage = data.get("usage", {})
            task = get_task(task_id)
            if task and "cost_metrics" in task:
                task["cost_metrics"]["llm_tokens_prompt"] += usage.get("prompt_tokens", 0)
                task["cost_metrics"]["llm_tokens_completion"] += usage.get("completion_tokens", 0)
                update_task(task_id, {"cost_metrics": task["cost_metrics"]})
            
            print(f"[Appendix Agent] Successfully generated Appendix ({len(appendix_text)} chars)")
            return appendix_text.strip()
        else:
            print(f"[Appendix Agent] Failed with HTTP status {response.status_code}")
            return ""
    except Exception as e:
        print(f"[Appendix Agent] Error: {e}")
        return ""



# -----------------------------
# Image Placeholder Post-Processing (Hybrid B Strategy)
# -----------------------------
def replace_image_placeholders(markdown_text: str, 
                                extraction_payload: Dict[str, Any]) -> str:
    """Convert AI-generated <<IMAGE:filename>> placeholders to real Pandoc-compatible markdown.
    
    This is the Hybrid B strategy:
    - AI decides WHICH images to use and WHERE to place them
    - This function does the actual markdown injection with proper paths
    
    Rules:
    - Only embeds images that exist in extraction_payload
    - Each image embedded only ONCE across entire document (deduplication)
    - Unknown placeholders are silently removed
    
    Args:
        markdown_text: Raw markdown from writer agent (may contain <<IMAGE:xxx>>)
        extraction_payload: The merged extraction data with image info
    
    Returns:
        Markdown with placeholders replaced by `images/filename` references
    """
    print(f"\n[ImageInject] ===== STARTING IMAGE PLACEHOLDER REPLACEMENT =====")
    
    # Step 1: Build a set of valid filenames from extraction
    valid_filenames = set()
    for slide in extraction_payload.get("slides", []):
        for img in slide.get("images", []):
            filename = img.get("filename", "")
            if filename:
                valid_filenames.add(filename)
    
    print(f"[ImageInject] Total valid image filenames: {len(valid_filenames)}")
    
    # Step 2: Track which images have already been used (deduplication)
    used_filenames = set()
    
    # Step 3: Find and replace all placeholders
    # Pattern matches: <<IMAGE:filename.ext>>
    placeholder_pattern = re.compile(r'<<IMAGE:([^>]+)>>')
    
    def replace_placeholder(match):
        filename = match.group(1).strip()
        
        # Case A: Unknown filename → remove
        if filename not in valid_filenames:
            print(f"[ImageInject]   ⚠️  Unknown filename, removing: {filename}")
            return ""
        
        # Case B: Already used → remove (no duplicates)
        if filename in used_filenames:
            print(f"[ImageInject]   ⚠️  Duplicate image, removing: {filename}")
            return ""
        
        # Case C: Valid and unused → inject proper markdown
        used_filenames.add(filename)
        proper_path = f"images/{filename}"
        print(f"[ImageInject]   ✅ Injecting: {filename}")
        return f"![]({proper_path})"
    
    fixed_markdown = placeholder_pattern.sub(replace_placeholder, markdown_text)
    
    # Step 4: Final stats
    print(f"[ImageInject] Total images injected: {len(used_filenames)}")
    print(f"[ImageInject] ====================================\n")
    
    return fixed_markdown



# =============================================================================
# Design Document Title Page & Front Matter
# =============================================================================

def _build_title_page(doc, first_p, project_title, document_title, client_name, client_logo_path):
    """Build the cover/title page at the very front of the document."""
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    first_p.paragraph_format.page_break_before = False

    def set_p_bottom_border(p, sz="144"):
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), sz)
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        pBdr.append(bottom)

    # Thick bar
    p_bar = first_p.insert_paragraph_before("")
    set_p_bottom_border(p_bar)

    # Project title
    p_title = first_p.insert_paragraph_before()
    p_title.paragraph_format.space_before = Pt(6)
    run = p_title.add_run(project_title)
    run.font.name = "Segoe UI"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # Subtitle
    p_sub = first_p.insert_paragraph_before()
    p_sub.paragraph_format.space_before = Pt(24)
    run = p_sub.add_run("Planning & Design Document")
    run.font.name = "Segoe UI"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # Client logo
    if client_logo_path and os.path.exists(client_logo_path):
        p_logo = first_p.insert_paragraph_before()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_logo.paragraph_format.space_before = Pt(40)
        run = p_logo.add_run()
        try:
            run.add_picture(str(client_logo_path), height=Inches(1.5))
        except Exception as e:
            print(f"[build_docx] Failed to insert client logo on title page: {e}")

    # Spacing
    for _ in range(4):
        first_p.insert_paragraph_before("")

    # Thin divider
    p_abs_div = first_p.insert_paragraph_before("")
    set_p_bottom_border(p_abs_div, sz="8")

    # Abstract
    p_abs_title = first_p.insert_paragraph_before()
    p_abs_title.paragraph_format.space_before = Pt(6)
    run = p_abs_title.add_run("Abstract")
    run.font.name = "Segoe UI"
    run.bold = True
    run.font.size = Pt(11)

    p_abs_text = first_p.insert_paragraph_before()
    p_abs_text.paragraph_format.space_before = Pt(6)
    run = p_abs_text.add_run(f"This document defines the Planning and Design for {document_title} project.")
    run.font.name = "Segoe UI"
    run.font.size = Pt(11)

    for _ in range(2):
        first_p.insert_paragraph_before("")

    # Prepared By
    p_prep = first_p.insert_paragraph_before()
    run = p_prep.add_run("Prepared By")
    run.font.name = "Segoe UI"
    run.font.size = Pt(11)

    enfrasys_logo = SCRIPT_DIR / "enfrasys-logo.jpg"
    if os.path.exists(enfrasys_logo):
        p_enfrasys = first_p.insert_paragraph_before()
        p_enfrasys.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_enfrasys.paragraph_format.space_before = Pt(10)
        run = p_enfrasys.add_run()
        try:
            run.add_picture(str(enfrasys_logo), width=Inches(1.5))
        except Exception as e:
            print(f"[build_docx] Failed to insert Enfrasys logo on title page: {e}")

    # Page break after title page
    p_break_title = first_p.insert_paragraph_before("")
    p_break_title.add_run().add_break(WD_BREAK.PAGE)


def _build_front_matter_and_toc(doc, first_p, client_name, project_title):
    """Build Change Record, Distribution List, Legal, TOC, and Sign-Off pages."""
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    anchor_p = first_p.insert_paragraph_before("")
    p_xml = anchor_p._element

    def add_p_before(text="", bold=False, size=None, space_before=None):
        p = anchor_p.insert_paragraph_before()
        if space_before is not None:
            p.paragraph_format.space_before = Pt(space_before)
        if text:
            r = p.add_run(text)
            if bold: r.bold = True
            if size: r.font.size = Pt(size)
            r.font.name = 'Segoe UI'
        return p

    # --- PAGE 2: Change Record & Distribution ---
    add_p_before("Change Record", bold=True, size=14)
    table1 = doc.add_table(rows=3, cols=4)
    try: table1.style = 'Table Grid'
    except KeyError as e: print(f"[build_docx] 'Table Grid' style not found for Change Record table: {e}")
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = 'Date'
    hdr_cells[1].text = 'Author'
    hdr_cells[2].text = 'Version'
    hdr_cells[3].text = 'Change Reference'
    row_cells = table1.rows[1].cells
    row_cells[0].text = '[change-record-date]'
    row_cells[1].text = '[change-record-author]'
    row_cells[2].text = '[change-record-version]'
    row_cells[3].text = '[change-record-reference]'
    p_xml.addprevious(table1._element)

    add_p_before("", space_before=12)
    add_p_before("Distribution List", bold=True, size=14)
    table2 = doc.add_table(rows=3, cols=2)
    try: table2.style = 'Table Grid'
    except KeyError as e: print(f"[build_docx] 'Table Grid' style not found for Distribution List table: {e}")
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = 'Name'
    hdr_cells2[1].text = 'Position/ Team'
    row_cells2 = table2.rows[1].cells
    row_cells2[0].text = '[Client/Entity/Company Name]'
    row_cells2[1].text = '[Client/Entity/Company Position]'
    row_cells3 = table2.rows[2].cells
    row_cells3[0].text = ''
    row_cells3[1].text = ''
    p_xml.addprevious(table2._element)

    add_p_before("", space_before=24)

    def add_legal(title, text):
        p = add_p_before(title, bold=True, size=7, space_before=12)
        p.paragraph_format.left_indent = Inches(2.3)
        ptext = add_p_before(text, size=7)
        ptext.paragraph_format.left_indent = Inches(2.3)
        ptext.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_legal("Terms of Use", "This work may be used \"as-is\" by any interested party. You may copy, adapt, and redistribute this document for non-commercial use or for your own internal use in a commercial setting. However, you may not republish this document, nor may you publish or distribute any adaptation of this document for other than non-commercial use or your own internal use, without first obtaining express written approval from Enfrasys Solutions Sdn Bhd.")
    add_legal("Disclaimer", "The Author and Enfrasys Solutions Sdn Bhd shall have neither liability nor responsibility to any person or entity with respect to the loss or damages arising from the information contained in this work. This work may include inaccuracies or typographical errors and solely represent the opinions of the Author. Changes are periodically made to this document without notice.Due to the rapid growth of various technologies, the Author and Enfrasys Solutions Sdn Bhd cannot guarantee the accuracy of the information presented after the date of publication.")
    add_legal("Trademarks", "The names of actual companies, service marks, trademarks or products mentioned herein may be the trademarks of their respective owners. Use of terms within this work should not be regarded as affecting the validity of any trademark or service mark. Enfrasys Consulting may have patents, patent applications, trademarks, copyrights, or other intellectual property rights covering subject matter in this document. Except as expressly provided in any written license agreement from Enfrasys Solutions Sdn Bhd, the furnishing of this document does not give you any license to those items.")

    p_break2 = add_p_before("")
    p_break2.add_run().add_break(WD_BREAK.PAGE)

    # --- PAGE 3: TABLE OF CONTENTS (XML INJECTION) ---
    p_toc_title = add_p_before("Table of Contents", size=15, bold=True)
    p_toc_title.paragraph_format.space_after = Pt(12)

    p_toc = add_p_before("")
    run_toc = p_toc.add_run()

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    run_toc._r.append(fldChar1)
    run_toc._r.append(instrText)
    run_toc._r.append(fldChar2)
    run_toc._r.append(fldChar3)

    p_break_toc = add_p_before("")
    p_break_toc.add_run().add_break(WD_BREAK.PAGE)

    # Force Word/LibreOffice to update fields on open
    doc_settings = doc.settings.element
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    doc_settings.append(update_fields)

    # --- PAGE 4: SIGN OFF ---
    p_signoff = add_p_before("1.0 Design Document Sign Off", size=15)
    try: p_signoff.style = 'Heading 1'
    except Exception as e: print(f"[build_docx] Failed to set Sign Off heading style: {e}")

    signoff_text_1 = f"We hereby acknowledge that the design document for {client_name} {project_title} has been reviewed, and all key aspects have been addressed satisfactorily."
    add_p_before(signoff_text_1, size=11, space_before=12)
    signoff_text_2 = "This document has been prepared, reviewed, accepted, and signed off by the following individuals:"
    add_p_before(signoff_text_2, size=11, space_before=12)
    add_p_before("", space_before=4)

    def set_table_borders(table):
        tbl_borders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '8')
            border.set(qn('w:color'), '000000')
            tbl_borders.append(border)
        table._tbl.tblPr.append(tbl_borders)

    # Signature table 1
    table_sig_1 = doc.add_table(rows=1, cols=2)
    try: table_sig_1.style = 'Table Grid'
    except Exception as e: print(f"[build_docx] 'Table Grid' style not found for signature table 1: {e}")
    set_table_borders(table_sig_1)
    s1_cells = table_sig_1.rows[0].cells
    p_s1_left = s1_cells[0].paragraphs[0]
    p_s1_left.add_run("Prepared by:\n\n")
    p_s1_left.add_run("Enfrasys Solutions Sdn Bhd").bold = True
    p_s1_left.add_run("\n\nSignature:\n\n\n\n\n___________________________\nName:\nDesignation:\nDate:")
    p_s1_right = s1_cells[1].paragraphs[0]
    p_s1_right.add_run("Verified by:\n\n")
    p_s1_right.add_run("Enfrasys Solutions Sdn Bhd").bold = True
    p_s1_right.add_run("\n\nSignature:\n\n\n\n\n___________________________\nName:\nDesignation:\nDate:")
    p_xml.addprevious(table_sig_1._element)

    p_break3 = add_p_before("")
    p_break3.add_run().add_break(WD_BREAK.PAGE)

    # Signature table 2
    table_sig_2 = doc.add_table(rows=4, cols=2)
    try: table_sig_2.style = 'Table Grid'
    except Exception as e: print(f"[build_docx] 'Table Grid' style not found for signature table 2: {e}")
    set_table_borders(table_sig_2)
    r1_cells = table_sig_2.rows[0].cells
    r1_cells[0].merge(r1_cells[1])
    r1_cells[0].text = ""
    p_r1 = r1_cells[0].paragraphs[0]
    p_r1.add_run("Reviewed by:\n").bold = True
    p_r1.add_run("\n\n")
    r2_cells = table_sig_2.rows[1].cells
    r2_cells[0].text = ""
    r2_cells[0].paragraphs[0].add_run("Name").bold = True
    r2_cells[1].text = ""
    r2_cells[1].paragraphs[0].add_run("Signature").bold = True
    r3_cells = table_sig_2.rows[2].cells
    r3_cells[0].text = "Name:\n\nDesignation:\n\nDate:"
    r3_cells[1].text = "\n\n\n\n\n___________________________"
    r4_cells = table_sig_2.rows[3].cells
    r4_cells[0].text = ""
    p_r4_L = r4_cells[0].paragraphs[0]
    p_r4_L.add_run("Verified by:\n\n").bold = True
    p_r4_L.add_run(f"[Designation]\n\n{client_name}\n\nSignature:\n\n\n\n\n___________________________\nName:\n\nDesignation:\n\nDate:")
    r4_cells[1].text = ""
    p_r4_R = r4_cells[1].paragraphs[0]
    p_r4_R.add_run("Approved by:\n\n").bold = True
    p_r4_R.add_run(f"[Designation]\n\n{client_name}\n\nSignature:\n\n\n\n\n___________________________\nName:\n\nDesignation:\n\nDate:")
    p_xml.addprevious(table_sig_2._element)

    p_break4 = add_p_before("")
    p_break4.add_run().add_break(WD_BREAK.PAGE)

    # Remove the anchor paragraph
    p_elem_anchor = anchor_p._element
    p_elem_anchor.getparent().remove(p_elem_anchor)


# =============================================================================
# Master DOCX Builder (Design Document)
# =============================================================================

def convert_md_to_docx(md_path, docx_path, document_title, project_title, client_name, client_logo_path=None):
    """Master function: Markdown → DOCX with full Enfrasys formatting.
    
    Pipeline:
    1. Pandoc converts MD → raw DOCX (with template styles)
    2. python-docx post-processes: title page, headers, footers, front matter, TOC, sign-off
    3. Enterprise table fixes (explicit XML borders, fixed layout, header banding)
    4. Typography enforcement (fonts, captions, page breaks)
    5. LibreOffice headless updates TOC fields
    """
    try:
        from docx import Document

        # Phase 1: Pandoc conversion
        _pandoc_md_to_raw_docx(str(md_path), str(docx_path))

        # Phase 2: Load and post-process
        doc = Document(str(docx_path))

        # Prepare first paragraph for front-page insertion
        if len(doc.paragraphs) > 0:
            first_p = doc.paragraphs[0]
            p_element = first_p._p
            for child in list(p_element):
                p_element.remove(child)
        else:
            first_p = doc.add_paragraph()

        # Build document structure
        _build_title_page(doc, first_p, project_title, document_title, client_name, client_logo_path)
        _apply_header_footer(doc, document_title, client_logo_path)
        _build_front_matter_and_toc(doc, first_p, client_name, project_title)

        # Phase 3: Fix content formatting
        _fix_images(doc)
        _fix_tables(doc)
        _fix_typography_and_captions(doc)
        _strip_heading_numbers(doc)

        # Save intermediate DOCX
        doc.save(str(docx_path))

        # Phase 4: Update TOC fields via LibreOffice
        try:
            _update_toc_with_libreoffice(str(docx_path))
            print("[convert_md_to_docx] TOC updated via LibreOffice roundtrip.")
        except Exception as e:
            print(f"[convert_md_to_docx] TOC roundtrip failed/skipped: {e}")

    except Exception as e:
        print(f"[convert_md_to_docx] Warning: {e}")
        traceback.print_exc()


# =============================================================================
# Main Background Processing Pipeline (Design Document)
# =============================================================================

def background_processing(task_id: str):
    task = get_task(task_id)
    if not task:
        return

    file_path = task["file_path"]
    logo_path = task.get("logo_path")
    task_dir = UPLOAD_DIR / task_id

    update_task(task_id, {
        "status": "processing",
        "step_name": "Extracting Structure & Images",
        "progress": 15
    })

    try:
        working_pptx_path = str(task_dir / "working.pptx")
        working_pdf_path = str(task_dir / "working.pdf")

        shutil.copy2(file_path, working_pptx_path)
        pptx_data = extract_with_pptx(working_pptx_path, task_dir)
        
        # Override presenter notes if user provided custom context
        user_notes_path = task_dir / "user_notes.json"
        if user_notes_path.exists():
            with open(user_notes_path, "r", encoding="utf-8") as f:
                user_notes = json.load(f)
            for slide in pptx_data.get("slides", []):
                s_num = str(slide["slide_number"])
                if s_num in user_notes:
                    slide["notes"] = user_notes[s_num]

        update_task(task_id, {
            "step_name": "Converting PPTX to PDF & AI Analysis",
            "progress": 30
        })

        convert_pptx_to_pdf(working_pptx_path, working_pdf_path)
        cu_data = analyze_with_content_understanding(working_pdf_path)

        update_task(task_id, {
            "step_name": "Analyzing Images with GPT-4o Vision",
            "progress": 50
        })

        pptx_data["slides"] = analyze_images_with_vision(task_id, pptx_data["slides"])

        merged = merge_extraction_results(pptx_data, cu_data)
        payload_file = task_dir / "extraction_payload.json"
        with open(payload_file, "w", encoding="utf-8") as f:
            json.dump(merged, f, indent=2, ensure_ascii=False, default=str)

        update_task(task_id, {
            "step_name": "Orchestrating Document Structure",
            "progress": 70
        })

        toc = generate_table_of_contents(task_id, merged)
        if not toc:
            raise Exception("Failed to generate Table of Contents.")

        update_task(task_id, {
            "step_name": "Drafting Document Sections (Taking 2-3 mins)",
            "progress": 85
        })

        # =============================================================================
        final_doc_md = write_document_sections(task_id, toc, merged)
        final_doc_md = final_doc_md.replace(str(task_dir) + os.sep, "")

        # NEW: Replace <<IMAGE:filename>> placeholders with proper markdown
        final_doc_md = replace_image_placeholders(final_doc_md, merged)

        # Step: Generate Appendix from the completed document (dedicated agent)
        update_task(task_id, {
            "step_name": "Generating Appendix References",
            "progress": 92
        })

        # Determine the last section number from the TOC
        last_section = 1
        for section in toc.get("sections", []):
            sec_num_str = section.get("section_number", "0.0")
            try:
                main_num = int(sec_num_str.split(".")[0])
                if main_num > last_section:
                    last_section = main_num
            except (ValueError, AttributeError) as e:
                print(f"[background_processing] Failed to parse section number '{sec_num_str}': {e}")

        print(f"[background_processing] Generating Appendix as section {last_section + 1}.0")
        appendix_md = generate_appendix_from_markdown(task_id, final_doc_md, last_section)

        if appendix_md:
            final_doc_md += f"\n\n{appendix_md}\n\n"
            print(f"[background_processing] Appendix appended successfully")
        else:
            print("[background_processing] Appendix generation skipped or failed; continuing without Appendix.")

        md_file = task_dir / "FINAL_DESIGN_DOCUMENT.md"
        with open(md_file, "w", encoding="utf-8") as f:
            f.write(final_doc_md)
        # =============================================================================

        update_task(task_id, {
            "step_name": "Converting to Word Document",
            "progress": 95
        })

        docx_path = task_dir / "Solution_Design_Document.docx"
        client_name = toc.get("client_name", "").strip()
        project_title = toc.get("project_title", "Project").strip()
        doc_title = f"{client_name} {project_title}" if client_name else project_title

        convert_md_to_docx(md_file, docx_path, doc_title, project_title, client_name, logo_path)

        # Final Cost Metrics (record content understanding pages)
        final_task = get_task(task_id)
        if final_task and "cost_metrics" in final_task:
            metrics = final_task["cost_metrics"]
            metrics["content_understanding_pages"] = len(merged.get("slides", []))
            update_task(task_id, {"cost_metrics": metrics})

        # Upload to blob storage
        blob_res = {"uploaded": False, "files": {}}
        try:
            blob_res = upload_task_to_blob(task_id, task_dir)
        except Exception as e:
            print(f"[background_processing] Blob upload failed (non-critical): {e}")

        # Calculate final filename
        raw_name = task.get("filename", "source.pptx")
        if raw_name.startswith("source_"):
            raw_name = raw_name[7:]
        if raw_name.lower().endswith(".pptx"):
            raw_name = raw_name[:-5]
        safe_name = re.sub(r'[\s-]+', '_', raw_name)
        generated_filename = f"generated_{safe_name}"

        docx_blob_name = (blob_res.get("files") or {}).get("docx") or f"{task_id}/Solution_Design_Document.docx"

        update_task(task_id, {
            "status": "completed",
            "step_name": "Ready!",
            "progress": 100,
            "generated_filename": generated_filename,
            "blobPath": docx_blob_name,
            "storage": blob_res
        })

    except Exception as e:
        error_trace = traceback.format_exc()
        print(f"[background_processing] FAILED for task {task_id}:")
        print(error_trace)
        update_task(task_id, {
            "status": "failed",
            "step_name": "Failed",
            "progress": 0,
            "error": str(e),
            "error_trace": error_trace
        })
