"""
final_config_docx.py — DOCX builder for the Final Configuration Document.

Kept separate from design_doc_agents.py so the Design Document builder stays untouched.

Front matter (agreed Phase 2 scope):
- Title page: project title + "Final Configuration", client logo, abstract, Prepared By
- Change Record + Distribution List (placeholders), Terms of Use / Disclaimer / Trademarks
- Table of Contents ONLY (no Table of Figures, no Table of Tables)
- NO sign-off page, NO executive summary, NO appendix
- Footer: "{Client} | {Project} | Final Configuration" + page number
"""

import os
import traceback

from config import SCRIPT_DIR
from doc_styles import (
    _pandoc_md_to_raw_docx,
    _apply_header_footer,
    _fix_tables,
    _fix_typography_and_captions,
    _fix_images,
    _strip_heading_numbers,
    _update_toc_with_libreoffice,
)

DOCUMENT_TYPE_LABEL = "Final Configuration"

LEGAL_SECTIONS = [
    ("Terms of Use",
     "This work may be used \"as-is\" by any interested party. You may copy, adapt, and redistribute this document "
     "for non-commercial use or for your own internal use in a commercial setting. However, you may not republish "
     "this document, nor may you publish or distribute any adaptation of this document for other than non-commercial "
     "use or your own internal use, without first obtaining express written approval from Enfrasys Solutions Sdn Bhd."),
    ("Disclaimer",
     "The Author and Enfrasys Solutions Sdn Bhd shall have neither liability nor responsibility to any person or entity "
     "with respect to the loss or damages arising from the information contained in this work. This work may include "
     "inaccuracies or typographical errors and solely represent the opinions of the Author. Changes are periodically "
     "made to this document without notice. Due to the rapid growth of various technologies, the Author and Enfrasys "
     "Solutions Sdn Bhd cannot guarantee the accuracy of the information presented after the date of publication."),
    ("Trademarks",
     "The names of actual companies, service marks, trademarks or products mentioned herein may be the trademarks of "
     "their respective owners. Use of terms within this work should not be regarded as affecting the validity of any "
     "trademark or service mark. Enfrasys Consulting may have patents, patent applications, trademarks, copyrights, or "
     "other intellectual property rights covering subject matter in this document. Except as expressly provided in any "
     "written license agreement from Enfrasys Solutions Sdn Bhd, the furnishing of this document does not give you any "
     "license to those items."),
]


# -----------------------------
# Small helpers
# -----------------------------
def _set_bottom_border(p, sz="144"):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_text(anchor, text="", size=11, bold=False, space_before=None):
    """Insert a paragraph directly before the anchor and return it."""
    from docx.shared import Pt, RGBColor

    p = anchor.insert_paragraph_before()
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if text:
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if bold:
            run.bold = bold
    return p


def _add_page_break(anchor):
    from docx.enum.text import WD_BREAK

    p = anchor.insert_paragraph_before("")
    p.add_run().add_break(WD_BREAK.PAGE)


def _add_picture(anchor, image_path, height=None, width=None, space_before=None):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if not image_path or not os.path.exists(image_path):
        return
    p = anchor.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    try:
        p.add_run().add_picture(str(image_path), height=height, width=width)
    except Exception as e:
        print(f"[final_config_docx] Failed to insert image {image_path}: {e}")


# -----------------------------
# Title page
# -----------------------------
def _build_fc_title_page(anchor, project_title, client_name, client_logo_path):
    from docx.shared import Inches

    # Thick bar
    p_bar = anchor.insert_paragraph_before("")
    _set_bottom_border(p_bar)

    # Client name (ALL CAPS) + project title + subtitle — matches the Serunai cover
    if client_name:
        _add_text(anchor, client_name.upper(), size=22, space_before=6)
    _add_text(anchor, project_title, size=22, bold=True, space_before=24 if client_name else 6)
    _add_text(anchor, DOCUMENT_TYPE_LABEL, size=22, space_before=24)

    # Client logo
    _add_picture(anchor, client_logo_path, height=Inches(1.5), space_before=40)

    for _ in range(4):
        _add_text(anchor)

    # Thin divider + Abstract
    p_div = anchor.insert_paragraph_before("")
    _set_bottom_border(p_div, sz="8")
    _add_text(anchor, "Abstract", bold=True, space_before=6)
    _add_text(
        anchor,
        f"This document defines the Final Configuration Document for {client_name} {project_title} project.",
        space_before=6,
    )

    for _ in range(2):
        _add_text(anchor)

    # Prepared By + Enfrasys logo
    _add_text(anchor, "Prepared By")
    _add_picture(anchor, SCRIPT_DIR / "enfrasys-logo.jpg", width=Inches(1.5), space_before=10)

    _add_page_break(anchor)


# -----------------------------
# Front matter (Change Record, Distribution, Legal, TOC) — NO sign-off
# -----------------------------
def _build_fc_front_matter(doc, anchor):
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def place_table(table):
        try:
            table.style = "Table Grid"
        except Exception as e:
            print(f"[final_config_docx] 'Table Grid' style not found: {e}")
        anchor._p.addprevious(table._element)

    # --- Change Record (placeholders) ---
    _add_text(anchor, "Change Record", size=14, bold=True)
    t1 = doc.add_table(rows=3, cols=4)
    for i, h in enumerate(["Date", "Author", "Version", "Change Reference"]):
        t1.rows[0].cells[i].text = h
    for i, v in enumerate(["[change-record-date]", "[change-record-author]",
                           "[change-record-version]", "[change-record-reference]"]):
        t1.rows[1].cells[i].text = v
    place_table(t1)

    # --- Distribution List (placeholders) ---
    _add_text(anchor, space_before=12)
    _add_text(anchor, "Distribution List", size=14, bold=True)
    t2 = doc.add_table(rows=3, cols=2)
    t2.rows[0].cells[0].text = "Name"
    t2.rows[0].cells[1].text = "Position/ Team"
    t2.rows[1].cells[0].text = "[Client/Entity/Company Name]"
    t2.rows[1].cells[1].text = "[Client/Entity/Company Position]"
    place_table(t2)

    # --- Legal text ---
    _add_text(anchor, space_before=24)
    for title, text in LEGAL_SECTIONS:
        p_title = _add_text(anchor, title, size=7, bold=True, space_before=12)
        p_title.paragraph_format.left_indent = Inches(2.3)
        p_text = _add_text(anchor, text, size=7)
        p_text.paragraph_format.left_indent = Inches(2.3)
        p_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _add_page_break(anchor)

    # --- Table of Contents (field) ---
    p_toc_title = _add_text(anchor, "Table of Contents", size=15, bold=True)
    p_toc_title.paragraph_format.space_after = Pt(12)

    p_toc = anchor.insert_paragraph_before("")
    run_toc = p_toc.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_end):
        run_toc._r.append(el)

    _add_page_break(anchor)

    # Ask Word to update fields (fills the TOC) when the document is opened
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)


# =============================================================================
# Master builder
# =============================================================================
def build_final_config_docx(md_path, docx_path, project_title, client_name, client_logo_path=None):
    """Markdown -> Final Configuration DOCX with Enfrasys formatting.

    Raises on failure so the pipeline marks the task as failed
    (instead of reporting 'completed' with a broken document).
    """
    from docx import Document

    try:
        # 1. Pandoc: Markdown -> raw DOCX (Enfrasys template styles)
        _pandoc_md_to_raw_docx(str(md_path), str(docx_path))
        doc = Document(str(docx_path))

        # 2. Insert a NEW empty anchor paragraph at the very top of the body.
        #    Unlike the Design Doc builder, we do NOT wipe the first paragraph,
        #    so no generated content is lost.
        anchor = doc.add_paragraph()
        doc.element.body.insert(0, anchor._p)

        # 3. Front pages (everything is inserted before the anchor, in order)
        _build_fc_title_page(anchor, project_title, client_name, client_logo_path)
        _apply_header_footer(
            doc,
            f"{client_name} | {project_title}",
            client_logo_path,
            document_type_label=DOCUMENT_TYPE_LABEL,
            title_separator=" | ",
        )
        _build_fc_front_matter(doc, anchor)

        # 4. Remove the anchor
        anchor._p.getparent().remove(anchor._p)

        # 5. Shared formatting pipeline
        _fix_images(doc)
        _fix_tables(doc)
        _fix_typography_and_captions(doc)

        # Table captions: centered, italic, grey (same look as Design Doc figure captions)
        import re
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor

        for p in doc.paragraphs:
            if re.match(r"^Table \d+ ", p.text.strip()):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(10)
                    run.font.italic = True
                    run.font.color.rgb = RGBColor(128, 128, 128)

        # The TOC page already ends with a page break, so the first section must not add another
        for p in doc.paragraphs:
            if p.style.name == "Heading 1":
                p.paragraph_format.page_break_before = False
                break

        _strip_heading_numbers(doc)
        doc.save(str(docx_path))

        # 6. Optional TOC pre-fill (non-critical; Word fills it on open anyway)
        try:
            _update_toc_with_libreoffice(str(docx_path))
            print("[final_config_docx] TOC updated via LibreOffice roundtrip.")
        except Exception as e:
            print(f"[final_config_docx] TOC roundtrip skipped (non-critical): {e}")

    except Exception as e:
        print(f"[final_config_docx] Failed to build DOCX: {e}")
        traceback.print_exc()
        raise